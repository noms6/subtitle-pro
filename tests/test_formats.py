"""TXT / DOCX / PDF pipelines and the CLI entry point."""
import sys

import pytest

import translate_srt as T
from conftest import SAMPLE_SRT, fake_translate

docx = pytest.importorskip("docx")


# ---------- TXT ----------
def test_text_file_roundtrip(engines, tmp_path):
    src = tmp_path / "notes.txt"
    src.write_bytes(b"Hello\n\nWorld\n")
    out = T.translate_text_file(str(src), "tr", 100, 4, "google")
    assert out.endswith("notes.tr.txt")
    raw = open(out, "rb").read()
    assert raw.startswith(b"\xef\xbb\xbf"), "output is written with a UTF-8 BOM"
    assert raw.decode("utf-8-sig") == f"{fake_translate('Hello')}\n\n{fake_translate('World')}\n"


@pytest.mark.xfail(strict=True, reason="H19: translate_text_file opens strictly as UTF-8; legacy Turkish (cp1254) .txt crashes")
def test_text_file_cp1254(engines, tmp_path):
    src = tmp_path / "eski.txt"
    src.write_bytes("Çok güzel\n".encode("cp1254"))
    T.translate_text_file(str(src), "tr", 100, 4, "google")


def test_text_file_overwrites_existing_output_without_asking(engines, tmp_path):
    src = tmp_path / "a.txt"
    src.write_text("Hi\n", encoding="utf-8")
    existing = tmp_path / "a.tr.txt"
    existing.write_text("my manual corrections", encoding="utf-8")
    T.translate_text_file(str(src), "tr", 100, 4, "google")
    assert "manual" not in existing.read_text(encoding="utf-8-sig")


# ---------- DOCX ----------
def _make_docx(path):
    d = docx.Document()
    p = d.add_paragraph("Hello ")
    p.add_run("brave").bold = True
    p.add_run(" new world.")
    t = d.add_table(rows=1, cols=2)
    t.cell(0, 0).text = "Cell text"
    t.cell(0, 1).text = "Another cell"
    d.sections[0].header.paragraphs[0].text = "Header line"
    # Hyperlink (w:hyperlink is not part of paragraph.runs in python-docx)
    from docx.oxml.shared import OxmlElement
    hp = d.add_paragraph("See: ")
    h = OxmlElement("w:hyperlink")
    r = OxmlElement("w:r")
    tx = OxmlElement("w:t")
    tx.text = "Click here"
    r.append(tx)
    h.append(r)
    hp._p.append(h)
    d.save(path)


def test_docx_body_table_header_translated(engines, tmp_path, monkeypatch):
    monkeypatch.setattr(T, "detect_language", lambda s: "en")
    src = tmp_path / "doc.docx"
    _make_docx(str(src))
    out = T.translate_docx_file(str(src), "tr", 100, 4, "google")
    d = docx.Document(out)
    full = [p.text for p in d.paragraphs]
    assert fake_translate("Cell text") == d.tables[0].cell(0, 0).text
    assert fake_translate("Header line") == d.sections[0].header.paragraphs[0].text
    assert any(fake_translate("brave") in t for t in full)


def test_docx_translates_run_fragments(engines, tmp_path, monkeypatch):
    """Characterises H25: a sentence split by bold formatting is sent as 3 fragments."""
    monkeypatch.setattr(T, "detect_language", lambda s: "en")
    src = tmp_path / "doc.docx"
    _make_docx(str(src))
    T.translate_docx_file(str(src), "tr", 100, 4, "google")
    sent = {t for _, t in engines.calls}
    assert {"Hello", "brave", "new world."} <= sent


@pytest.mark.xfail(strict=True, reason="H25: paragraphs should be translated as whole sentences, then mapped back onto runs")
def test_docx_paragraph_level(engines, tmp_path, monkeypatch):
    monkeypatch.setattr(T, "detect_language", lambda s: "en")
    src = tmp_path / "doc.docx"
    _make_docx(str(src))
    T.translate_docx_file(str(src), "tr", 100, 4, "google")
    assert ("google", "Hello brave new world.") in engines.calls


@pytest.mark.xfail(strict=True, reason="H25: hyperlink text (and text boxes / footnotes) is skipped")
def test_docx_hyperlink_text_translated(engines, tmp_path, monkeypatch):
    monkeypatch.setattr(T, "detect_language", lambda s: "en")
    src = tmp_path / "doc.docx"
    _make_docx(str(src))
    T.translate_docx_file(str(src), "tr", 100, 4, "google")
    assert ("google", "Click here") in engines.calls


def test_docx_same_language_is_copied(engines, tmp_path, monkeypatch):
    monkeypatch.setattr(T, "detect_language", lambda s: "tr")
    src = tmp_path / "doc.docx"
    _make_docx(str(src))
    out = T.translate_docx_file(str(src), "tr", 100, 4, "google")
    assert engines.calls == []
    assert out.endswith("doc.tr.docx")


def test_docx_corrupt_file_raises(engines, tmp_path):
    src = tmp_path / "bad.docx"
    src.write_bytes(b"not a zip")
    with pytest.raises(Exception):  # noqa: B017 - python-docx raises PackageNotFoundError
        T.translate_docx_file(str(src), "tr", 100, 4, "google")


# ---------- PDF ----------
reportlab = pytest.importorskip("reportlab")
pdfminer_hl = pytest.importorskip("pdfminer.high_level")


def _make_pdf(path, lines):
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas
    c = canvas.Canvas(path, pagesize=A4)
    y = 800
    for ln in lines:
        c.drawString(50, y, ln)
        y -= 16
    c.save()


def test_pdf_to_pdf(engines, tmp_path, monkeypatch):
    monkeypatch.setattr(T, "detect_language", lambda s: "en")
    src = tmp_path / "in.pdf"
    _make_pdf(str(src), ["Hello there.", "", "Second paragraph."])
    out = T.translate_pdf_file(str(src), "tr", 100, 4, "google")
    text = pdfminer_hl.extract_text(out)
    assert "TR<Hello there.>" in text


def test_pdf_turkish_glyphs_survive(tmp_path, monkeypatch):
    monkeypatch.setattr(T, "_build_translation_map", lambda texts, *a, **k: {t: "Çığ öğüş İŞ" for t in texts})
    src = tmp_path / "in.pdf"
    _make_pdf(str(src), ["Hello"])
    out = T.translate_pdf_file(str(src), "tr", 100, 4, "google")
    assert "Çığ öğüş İŞ" in pdfminer_hl.extract_text(out)


def test_pdf_as_text(engines, tmp_path, monkeypatch):
    monkeypatch.setattr(T, "detect_language", lambda s: "en")
    src = tmp_path / "in.pdf"
    _make_pdf(str(src), ["Hello there."])
    out = T.translate_pdf_file(str(src), "tr", 100, 4, "google", as_text=True)
    assert out.endswith("in.tr.txt")


@pytest.mark.xfail(strict=True, reason="H26: PDF hard line breaks split sentences; each visual line is translated separately")
def test_pdf_sentence_across_lines(engines, tmp_path, monkeypatch):
    monkeypatch.setattr(T, "detect_language", lambda s: "en")
    src = tmp_path / "in.pdf"
    _make_pdf(str(src), ["This sentence continues", "on the next line."])
    T.translate_pdf_file(str(src), "tr", 100, 4, "google")
    assert any("continues on the next" in t for _, t in engines.calls)


@pytest.mark.xfail(strict=True, reason="H26: long lines are wrapped by character count, cutting words in half")
def test_pdf_wrap_does_not_split_words(tmp_path, monkeypatch):
    long = " ".join(["internationalization"] * 12)
    monkeypatch.setattr(T, "_build_translation_map", lambda texts, *a, **k: {t: long for t in texts})
    src = tmp_path / "in.pdf"
    _make_pdf(str(src), ["x"])
    out = T.translate_pdf_file(str(src), "tr", 100, 4, "google")
    words = pdfminer_hl.extract_text(out).split()
    assert all(w == "internationalization" for w in words)


# ---------- CLI ----------
def test_cli_translates_file(engines, tmp_path, monkeypatch):
    src = tmp_path / "in.srt"
    src.write_bytes(SAMPLE_SRT.encode("utf-8"))
    dst = tmp_path / "out.srt"
    monkeypatch.setattr(sys, "argv", ["translate_srt.py", str(src), str(dst), "tr"])
    T._cli_main()
    out = dst.read_bytes()
    assert out.startswith(b"\xef\xbb\xbf")
    assert "TR<Stop! Stop the bus!>" in out.decode("utf-8-sig")
    assert b"\r\n" in out


@pytest.mark.xfail(strict=True, reason="H19: CLI reads input strictly as UTF-8 (GUI uses the robust reader)")
def test_cli_cp1254_input(engines, tmp_path, monkeypatch):
    src = tmp_path / "in.srt"
    src.write_bytes("1\r\n00:00:01,000 --> 00:00:02,000\r\nÇok güzel\r\n".encode("cp1254"))
    monkeypatch.setattr(sys, "argv", ["x", str(src), str(tmp_path / "o.srt"), "en"])
    T._cli_main()


@pytest.mark.xfail(strict=True, reason="H35: CLI exits 0 and writes the file even when every line failed to translate")
def test_cli_nonzero_exit_on_total_failure(engines, tmp_path, monkeypatch):
    engines.google_always_fail = True
    engines.mm_always_fail = True
    src = tmp_path / "in.srt"
    src.write_bytes(SAMPLE_SRT.encode("utf-8"))
    monkeypatch.setattr(sys, "argv", ["x", str(src), str(tmp_path / "o.srt"), "tr"])
    with pytest.raises(SystemExit) as ei:
        T._cli_main()
    assert ei.value.code != 0


def test_cli_missing_input_raises(engines, tmp_path, monkeypatch):
    monkeypatch.setattr(sys, "argv", ["x", str(tmp_path / "nope.srt"), str(tmp_path / "o.srt")])
    with pytest.raises(FileNotFoundError):
        T._cli_main()


def test_cli_has_no_help_flag(engines, monkeypatch):
    """H35: `--help` is treated as an input path (argparse not used)."""
    monkeypatch.setattr(sys, "argv", ["x", "--help", "out.srt"])
    with pytest.raises(FileNotFoundError):
        T._cli_main()
